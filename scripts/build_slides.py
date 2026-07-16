"""Generate the Supply Chain Atlas v2 slide deck (docs/Supply_Chain_Atlas.pptx).

The deck is deliberately minimal — 5 slides that match docs/VIDEO_SCRIPT.md
(title/hook, problem + product loop, architecture, verified numbers,
reflection). The video is mostly live UI demos; slides only frame them.
python-pptx is a doc-gen-only dependency, so run it ephemerally rather than
adding it to the project:

    uv run --with python-pptx python scripts/build_slides.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Supply_Chain_Atlas.pptx"

INK = RGBColor(0x16, 0x20, 0x2B)
MUTED = RGBColor(0x5C, 0x67, 0x77)
ACCENT = RGBColor(0x1F, 0x6F, 0xB8)
PANEL = RGBColor(0x0E, 0x14, 0x1B)
CODE = RGBColor(0xE6, 0xEC, 0xF3)
COMMENT = RGBColor(0x82, 0x92, 0xA8)
KEY = RGBColor(0x85, 0xB7, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _tb(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _run(p, text, size, color, *, bold=False, mono=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Consolas" if mono else "Segoe UI"
    return r


def _accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background(); bar.shadow.inherit = False


def _title(slide, title, note=None):
    _accent_bar(slide)
    tf = _tb(slide, 0.55, 0.35, 12.2, 1.1)
    _run(tf.paragraphs[0], title, 30, INK, bold=True)
    if note:
        p = tf.add_paragraph()
        _run(p, note, 15, MUTED)


def bullets_slide(title, bullets, note=None, *, top=1.7, size=16):
    s = prs.slides.add_slide(BLANK)
    _title(s, title, note)
    tf = _tb(s, 0.7, top, 12.0, 7.1 - top)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        lead, rest = (b.split("|", 1) + [""])[:2] if "|" in b else (None, b)
        _run(p, "▸ ", size, ACCENT, bold=True)
        if lead:
            _run(p, lead.strip() + " — ", size, INK, bold=True)
        _run(p, rest.strip(), size, INK)
    return s


def code_slide(title, code, note=None):
    s = prs.slides.add_slide(BLANK)
    _title(s, title, note)
    top = 1.55
    panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(top),
                               Inches(12.25), Inches(5.65))
    panel.fill.solid(); panel.fill.fore_color.rgb = PANEL
    panel.line.fill.background(); panel.shadow.inherit = False
    box = s.shapes.add_textbox(Inches(0.8), Inches(top + 0.15), Inches(11.8), Inches(5.35))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = code.strip("\n").split("\n")
    size = 13 if max((len(l) for l in lines), default=0) <= 66 and len(lines) <= 20 else 11
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05; p.space_after = Pt(0)
        stripped = line.lstrip()
        color = COMMENT if stripped.startswith(("#", "--", "/*", "*")) else CODE
        _run(p, line if line else " ", size, color, mono=True)


# ------------------------------------------------------------ slide 1: title
def title_slide():
    s = prs.slides.add_slide(BLANK)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.5), Inches(13.333), Inches(2.5))
    band.fill.solid(); band.fill.fore_color.rgb = PANEL
    band.line.fill.background(); band.shadow.inherit = False
    tf = _tb(s, 0.8, 2.75, 11.7, 2.0)
    _run(tf.paragraphs[0], "Supply Chain Atlas", 46, WHITE, bold=True)
    p = tf.add_paragraph()
    _run(p, "Detect supply-chain & ownership changes in SEC filings — alert, predict the "
            "market reaction, and measure the prediction",
         18, RGBColor(0xA5, 0xB2, 0xC2))
    p2 = tf.add_paragraph()
    _run(p2, "detect → alert → predict → measure  ·  Airflow 3 · dbt · Kafka/Flink · MLflow · FastAPI",
         14, KEY, mono=True)
    foot = _tb(s, 0.8, 6.6, 11.7, 0.5)
    _run(foot.paragraphs[0], "Data-engineering capstone · v2", 13, MUTED)


title_slide()

# ----------------------------------------------- slide 2: problem + the loop
s2 = bullets_slide("The problem — filings tell you who depends on whom", [
    "Scattered signal | insider trades, ownership stakes, subsidiaries, customer concentration — "
    "across a dozen filing types, ~2,000 filings a day, mostly messy text.",
    "The user | an analyst watching a supply-chain universe, who needs relationship CHANGES, "
    "not documents.",
    "Scope | a 79-entity hardware/semiconductor universe — and the universe itself is data "
    "(one CSV), so rescoping never touches code.",
], note="The user, the signal, and the product loop", top=1.6)
band = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(4.9),
                           Inches(12.0), Inches(1.7))
band.fill.solid(); band.fill.fore_color.rgb = PANEL
band.line.fill.background(); band.shadow.inherit = False
tf = _tb(s2, 1.0, 5.1, 11.4, 1.4)
_run(tf.paragraphs[0], "DETECT a relationship change  →  ALERT on it  →  PREDICT the market "
                       "reaction (3d / 1w / 1m / 3m)  →  MEASURE the prediction",
     16, CODE, mono=True)
p = tf.add_paragraph(); p.space_before = Pt(8)
_run(p, "# when the actual prices arrive, every prediction is scored — the loop closes daily",
     13, COMMENT, mono=True)

# -------------------------------------------------------- slide 3: architecture
code_slide("Architecture — batch backbone, event layer, ML loop, fast path", """
SOURCES  pluggable registry: EDGAR backfill + CDC poller . market bars (yfinance)
             |
BRONZE   gzipped docs on disk . Parquet lake (dt-partitioned)   # idempotent, cursor-resumable
             |
SILVER   4 fixture-tested parsers -> Postgres . price_daily
             |
GOLD     dbt star schema: conformed dims + 5 edge facts, SCD2 both flavors,
         incremental big facts, staging+marts tests -- 82 build steps
             |
EVENTS   change detectors -> filing_event  .  rule engine -> alert (dedup, severity)
             |
ML       stage 1: 8-K substance filter  .  stage 2: reaction model, 4 horizons
         predict @ event time -> backfill actuals -> 12:00 eval -> degradation alerts
         MLflow registry: @champion alias promotion, frozen-baseline PSI drift
             |
SERVING  FastAPI + live dashboard (graph . changes . outlook . pulse)

FAST PATH      EDGAR live feed -> Redpanda -> alert consumer -> same alert table
               trades -> Flink SQL (event time, watermarks, 10s VWAP) -> Postgres
ORCHESTRATION  Airflow 3 in Docker Compose (Asia/Singapore):
               daily 09:30 . eval 12:00 . retrain Sun 10:00 . NLP Sat 10:00
""", note="Bronze -> Silver -> Gold -> events -> ML -> serving; Kafka/Flink beside it; Airflow over all of it")

# ------------------------------------------------------ slide 4: verified numbers
bullets_slide("Results — every number verified against the running system", [
    "Data | 205,000+ filings + ~196,000 daily bars from 2 sources, over a 79-entity universe "
    "defined as data (one CSV).",
    "Warehouse | 82 green dbt build steps — staging + marts tests, SCD2 both flavors, "
    "incremental facts.",
    "Events → alerts | 6,085 real alerts across 5 event types, deduplicated and severity-graded.",
    "ML | 2 registered models behind @champion aliases; stage-1 filter AUC .989 / F1 .969.",
    "The loop | 24,340 predictions logged, 24,132 measured against realized abnormal returns; "
    "trailing hit rates 54–57% by horizon.",
    "The honest finding | stage-2 test AUC ~0.52 — near coin-flip, and the system says so, "
    "daily, automatically.",
    "Fast path | Kafka/Flink: 4,750+ event-time trade windows landed in Postgres.",
    "Operations | 4 Airflow DAGs (daily 09:30 · eval 12:00 · weekly retrain · weekly NLP, "
    "Asia/Singapore); survives a nightly 9-hour power-off.",
    "Quality | 64-test pytest suite — found 5 real parser bugs.",
], note="Counts from the live warehouse, registry, and Airflow — not projections", top=1.6, size=15)

# ---------------------------------------------------------- slide 5: reflection
bullets_slide("Reflection", [
    "Hardest | closing the actuals loop idempotently — event dates, trading-day windows, late "
    "bars, and a machine that sleeps 9 hours a night, all converging to the same state.",
    "Proudest | the system's honesty: the reaction model is near coin-flip and the system "
    "reports it daily, automatically — a dashboard that flatters you is worthless.",
    "With more time | curate the NLP gold set already gated in the weekly DAG; widen the "
    "universe to thousands of filers with concurrent ingestion.",
    "With more time | shadow-score challenger models on live events before promotion; bring "
    "in Spark when the universe goes whole-market.",
], note="Hardest part, proudest part, and what comes next")

n = len(list(prs.slides))
prs.save(str(OUT))
print(f"wrote {OUT} ({n} slides)")
